"""Data adapters for LIMS, ELN, and instrument software integration."""

from __future__ import annotations
from pydantic import BaseModel, Field
from pathlib import Path
from typing import Any
import csv
import json


class LIMSAdapter(BaseModel):
    name: str = "LIMS"
    api_url: str = ""
    api_key: str = ""


class InstrumentAdapter(BaseModel):
    name: str = "Instrument"
    file_patterns: list[str] = Field(default_factory=list)
    api_url: str = ""


class FileWatcherAdapter:
    """Monitor directories for new data files and parse them."""

    def __init__(self, watch_dirs: list[str], file_patterns: list[str] | None = None):
        self.watch_dirs = [Path(d) for d in watch_dirs]
        self.file_patterns = file_patterns or ["*.csv", "*.xlsx", "*.txt"]

    def scan_files(self) -> list[Path]:
        found = []
        for watch_dir in self.watch_dirs:
            if not watch_dir.exists():
                continue
            for pattern in self.file_patterns:
                found.extend(watch_dir.glob(pattern))
        return found

    def parse_csv(self, file_path: Path) -> list[dict]:
        records = []
        with open(file_path, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                records.append(dict(row))
        return records

    def parse_instrument_output(self, file_path: Path | str) -> dict:
        file_path = Path(file_path)
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            return {"records": self.parse_csv(file_path)}
        elif suffix == ".xlsx":
            return self._parse_xlsx(file_path)
        elif suffix == ".txt":
            return self._parse_text_report(file_path)
        else:
            return {"error": f"Unsupported file format: {suffix}"}

    def _parse_xlsx(self, file_path: Path) -> dict:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"error": "openpyxl is not installed; cannot parse .xlsx files"}
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return {"records": []}
        header = [str(c) if c is not None else "" for c in rows[0]]
        records = []
        for row in rows[1:]:
            records.append({header[i]: row[i] for i in range(len(header))})
        return {"records": records}

    def _parse_text_report(self, file_path: Path) -> dict:
        content = file_path.read_text(encoding="utf-8")
        lines = content.strip().split("\n")
        return {"lines": lines, "num_lines": len(lines)}
